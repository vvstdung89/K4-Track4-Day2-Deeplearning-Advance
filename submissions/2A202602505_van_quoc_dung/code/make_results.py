"""make_results.py - tổng hợp mọi lần chạy thành results.xlsx và các biểu đồ cho báo cáo (Bước 5).

    LAB_SUB=... LAB_DATA=... python make_results.py --backbone convnext_tiny

Mọi chỉ số test/val ở sheet Final và PerClass được TÍNH LẠI từ predictions/*.csv bằng
eval.read_pred + eval.compute_metrics (cùng định nghĩa với `python eval.py score`).
"""
from __future__ import annotations

import argparse
import glob
import json
import os
from pathlib import Path

import numpy as np
import pandas as pd

import dataset as D
import train as TR  # noqa: F401  (thêm repo root vào sys.path)
from eval import compute_metrics, mean_std, read_pred
from run_experiments import BACKBONES, TRAINING

SUB = Path(os.environ.get("LAB_SUB", Path(__file__).resolve().parent.parent))
DATA = Path(os.environ.get("LAB_DATA", "data"))
TABLES, FIGS = SUB / "tables", SUB / "figures"


def load_summaries() -> dict:
    out = {}
    for f in glob.glob(str(SUB / "runs" / "*" / "seed*" / "summary.json")):
        s = json.loads(Path(f).read_text(encoding="utf-8"))
        out[(s["exp_id"], s["seed"])] = s
    return out


def fmt_ms(m, s, nd=4):
    return f"{m:.{nd}f} ± {s:.{nd}f}" if np.isfinite(s) else f"{m:.{nd}f}"


def group_metrics(pattern: str) -> list[dict]:
    rows = []
    for f in sorted(glob.glob(pattern)):
        p = read_pred(f)
        m = compute_metrics(p.y_true, p.y_pred, p.probs)
        rows.append({"file": Path(f).name, "seed": p.seed, **m})
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--backbone", required=True)
    a = ap.parse_args()
    S = load_summaries()
    FIGS.mkdir(parents=True, exist_ok=True)
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    # ---------------- Backbones ----------------
    rows = []
    for eid, bb in BACKBONES:
        s = S.get((eid, 0))
        if not s:
            continue
        rows.append({"exp_id": eid, "backbone": bb, "weight_tag": s["weight_tag"], "params_M": s["params_M"],
                     "GMAC": s["gmacs"], "img_size": s["config"]["img_size"], "epochs": s["epochs"],
                     "seed": 0, "val_macro_f1": s["val_macro_f1"], "val_top1": s["val_top1"],
                     "val_f1_chinee": s["val_f1_per_class"][0], "val_f1_snake": s["val_f1_per_class"][7],
                     "best_epoch": s["best_epoch"], "train_time_per_epoch_s": s["train_time_per_epoch_s"],
                     "latency_b1_fp32_p50_ms": s["latency_b1_fp32_p50_ms"],
                     "latency_b1_fp32_p95_ms": s["latency_b1_fp32_p95_ms"],
                     "curve": f"curves/{eid}_{bb}.png",
                     "note": "công thức nền T00 (AdamW, LR 1e-4/1e-3, wd 0.05, warmup 1 + cosine, CE, AMP), 1 seed"})
    bdf = pd.DataFrame(rows)

    # ---------------- Training ----------------
    t00 = S.get(("T00", 0))
    t00_seeds = [S[("T00", k)]["val_macro_f1"] for k in (0, 1, 2) if ("T00", k) in S]
    noise = float(np.std(t00_seeds, ddof=1)) if len(t00_seeds) > 1 else float("nan")
    rows = []
    spec = [("T00", "-", "baseline", {})] + TRAINING + [("T10", "kết hợp", "combo", None)]
    for eid, axis, desc, ov in spec:
        s = S.get((eid, 0))
        if not s:
            continue
        if ov is None:
            base = dataclass_diff(t00["config"], s["config"]) if t00 else {}
            diff = ", ".join(f"{k}={v}" for k, v in base.items())
        else:
            diff = ", ".join(f"{k}={v}" for k, v in ov.items()) or "(nền)"
        rows.append({"exp_id": eid, "backbone": s["backbone"], "axis": axis, "diff_vs_T00": diff, "seed": 0,
                     "val_macro_f1": s["val_macro_f1"], "val_top1": s["val_top1"],
                     "delta_f1_vs_T00": s["val_macro_f1"] - t00["val_macro_f1"] if t00 else np.nan,
                     "delta_over_noise_std": (s["val_macro_f1"] - t00["val_macro_f1"]) / noise if t00 and noise > 0 else np.nan,
                     "val_f1_chinee": s["val_f1_per_class"][0], "val_f1_snake": s["val_f1_per_class"][7],
                     "val_ece": s["val_ece"], "best_epoch": s["best_epoch"],
                     "train_time_per_epoch_s": s["train_time_per_epoch_s"],
                     "curve": f"curves/{eid}_{s['desc']}.png",
                     "note": (f"= {s['alias_of']} (cùng cấu hình). " if s.get("alias_of") else "")
                             + (f"F1 trọng số thường (không EMA) tốt nhất {s['val_macro_f1_raw_best']:.4f}. "
                                if s.get("val_macro_f1_raw_best") else "")})
    tdf = pd.DataFrame(rows)

    # ---------------- Inference / Latency ----------------
    idf, ldf = read_table("inference.csv"), read_table("latency.csv")

    # ---------------- Final ----------------
    P = SUB / "predictions"
    groups = {
        "F01": (f"{a.backbone} + {combo_desc()} + {infer_desc()} + temperature scaling (T khớp trên val)",
                P / "F01_seed*_test.csv", P / "F01_seed*_val.csv", P / "F01uncal_seed*_test.csv"),
        "T00": (f"{a.backbone} + công thức nền T00 + 1-view I00 (mốc)", P / "T00_seed*_test.csv",
                P / "T00_seed*_val.csv", None),
    }
    frows, per_class, final_stats = [], {}, {}
    for gid, (desc, pt, pv, pu) in groups.items():
        test = group_metrics(str(pt))
        val = {r["seed"]: r for r in group_metrics(str(pv))}
        unc = {r["seed"]: r for r in group_metrics(str(pu))} if pu else {}
        for r in test:
            frows.append({"exp_id": gid, "config": desc, "seed": r["seed"],
                          "val_macro_f1": val.get(r["seed"], {}).get("macro_f1", np.nan),
                          "test_macro_f1": r["macro_f1"], "test_top1": r["top1"],
                          "test_balanced_acc": r["balanced_acc"], "test_ece": r["ece"],
                          "test_ece_uncalibrated": unc.get(r["seed"], {}).get("ece", np.nan),
                          "test_recall_chinee": r["recall"][0], "test_recall_snake": r["recall"][7],
                          "test_f1_chinee": r["f1"][0], "test_f1_snake": r["f1"][7],
                          "pred_file": r["file"]})
        if test:
            st = {}
            for k in ("macro_f1", "top1", "balanced_acc", "ece"):
                st[k] = mean_std([r[k] for r in test])
            st["val_macro_f1"] = mean_std([val[r["seed"]]["macro_f1"] for r in test if r["seed"] in val])
            st["recall_chinee"] = mean_std([r["recall"][0] for r in test])
            st["recall_snake"] = mean_std([r["recall"][7] for r in test])
            if unc:
                st["ece_uncal"] = mean_std([unc[r["seed"]]["ece"] for r in test if r["seed"] in unc])
            final_stats[gid] = st
            frows.append({"exp_id": f"{gid} (mean ± std, {len(test)} seed)", "config": desc, "seed": "all",
                          "val_macro_f1": fmt_ms(*st["val_macro_f1"]), "test_macro_f1": fmt_ms(*st["macro_f1"]),
                          "test_top1": fmt_ms(*st["top1"]), "test_balanced_acc": fmt_ms(*st["balanced_acc"]),
                          "test_ece": fmt_ms(*st["ece"]),
                          "test_ece_uncalibrated": fmt_ms(*st["ece_uncal"]) if "ece_uncal" in st else "",
                          "test_recall_chinee": fmt_ms(*st["recall_chinee"]),
                          "test_recall_snake": fmt_ms(*st["recall_snake"])})
            pcm = {k: mean_std([r[k] for r in test]) for k in ("precision", "recall", "f1")}
            per_class[gid] = (pcm, test[0]["support"], test)
    fdf = pd.DataFrame(frows)
    if "F01" in final_stats and "T00" in final_stats:
        d = final_stats["F01"]["macro_f1"][0] - final_stats["T00"]["macro_f1"][0]
        s = max(final_stats["F01"]["macro_f1"][1], final_stats["T00"]["macro_f1"][1])
        fdf = pd.concat([fdf, pd.DataFrame([{"exp_id": "Δ F01 − T00", "config": "chênh lệch macro-F1 test (mean)",
                                             "test_macro_f1": f"{d:+.4f} (std lớn hơn = {s:.4f}; Δ/std = {d / s:.1f})"}])])

    # ---------------- PerClass ----------------
    prow = []
    for gid, (pcm, support, _) in per_class.items():
        for c in range(9):
            prow.append({"config": gid, "class": D.CLASS_NAMES[c], "n_test": int(support[c]),
                         "precision": fmt_ms(pcm["precision"][0][c], pcm["precision"][1][c]),
                         "recall": fmt_ms(pcm["recall"][0][c], pcm["recall"][1][c]),
                         "f1": fmt_ms(pcm["f1"][0][c], pcm["f1"][1][c])})
    pdf = pd.DataFrame(prow)

    # ---------------- Summary ----------------
    srows = []
    for _, r in bdf.iterrows():
        srows.append({"exp_id": r.exp_id, "stage": "Backbone", "description": r.backbone,
                      "val_macro_f1": r.val_macro_f1, "val_top1": r.val_top1, "GMAC": r.GMAC,
                      "lat_b1_p50_ms": r.latency_b1_fp32_p50_ms, "seeds": 1})
    for _, r in tdf.iterrows():
        if r.exp_id == "T00":
            continue
        srows.append({"exp_id": r.exp_id, "stage": "Training", "description": f"{r.backbone}: {r.diff_vs_T00}",
                      "val_macro_f1": r.val_macro_f1, "val_top1": r.val_top1, "seeds": 1})
    for _, r in idf.iterrows():
        if r.exp_id == "I00":
            continue
        srows.append({"exp_id": r.exp_id, "stage": "Inference", "description": r.method,
                      "val_macro_f1": r.val_macro_f1, "val_top1": r.val_top1,
                      "lat_b1_p50_ms": r.get("lat_p50_ms", np.nan), "seeds": 1})
    sdf = pd.DataFrame(srows)
    if len(sdf):
        sdf = sdf.sort_values("val_macro_f1", ascending=False).head(10)
    head = []
    for gid in ("F01", "T00"):
        if gid in final_stats:
            st = final_stats[gid]
            head.append({"exp_id": gid, "stage": "FINAL (test)" if gid == "F01" else "BASELINE (test)",
                         "description": groups[gid][0], "val_macro_f1": fmt_ms(*st["val_macro_f1"]),
                         "test_macro_f1": fmt_ms(*st["macro_f1"]), "test_top1": fmt_ms(*st["top1"]),
                         "test_ece": fmt_ms(*st["ece"]), "seeds": 3})
    summary = pd.concat([pd.DataFrame(head), pd.DataFrame([{}]), sdf], ignore_index=True)

    out = SUB / "results.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in (("Summary", summary), ("Backbones", bdf), ("Training", tdf), ("Inference", idf),
                         ("Final", fdf), ("PerClass", pdf), ("Latency", ldf)):
            df.to_excel(xw, sheet_name=name, index=False)
    style_workbook(out)
    print("đã ghi", out)

    # ---------------- Figures ----------------
    if len(bdf) and bdf.latency_b1_fp32_p50_ms.notna().all():
        fig, ax = plt.subplots(1, 2, figsize=(11, 4))
        for _, r in bdf.iterrows():
            ax[0].scatter(r.latency_b1_fp32_p50_ms, r.val_macro_f1, s=40 + 4 * r.params_M)
            ax[0].annotate(r.backbone, (r.latency_b1_fp32_p50_ms, r.val_macro_f1), fontsize=8)
            ax[1].scatter(r.GMAC, r.train_time_per_epoch_s)
            ax[1].annotate(r.backbone, (r.GMAC, r.train_time_per_epoch_s), fontsize=8)
        ax[0].set_xlabel("độ trễ batch 1 FP32 p50 (ms, T4)"); ax[0].set_ylabel("macro-F1 val")
        ax[0].set_title("Backbone: chất lượng vs độ trễ (cỡ điểm ∝ #params)")
        ax[1].set_xlabel("GMAC / ảnh"); ax[1].set_ylabel("thời gian train / epoch (s)")
        ax[1].set_title("FLOPs có dự đoán thời gian train?")
        for x in ax: x.grid(alpha=.3)
        fig.tight_layout(); fig.savefig(FIGS / "backbones_tradeoff.png", dpi=110); plt.close(fig)
    if len(tdf) > 1:
        # T01/T02 (khởi tạo) lệch hàng chục lần các trục khác: vẽ riêng để thấy được chênh lệch nhỏ
        t = tdf[~tdf.exp_id.isin(["T00", "T01", "T02"])]
        fig, ax = plt.subplots(figsize=(9, 4.2))
        ax.barh([f"{r.exp_id} {r.diff_vs_T00[:40]}" for _, r in t.iterrows()], t.delta_f1_vs_T00,
                color=["tab:green" if v > 0 else "tab:red" for v in t.delta_f1_vs_T00])
        if np.isfinite(noise):
            d = noise * np.sqrt(2)
            ax.axvspan(-d, d, color="gray", alpha=.25, label=f"± std của hiệu 2 lần chạy 1 seed (√2·{noise:.4f})")
            ax.legend(loc="lower right", fontsize=8)
        ax.axvline(0, color="k", lw=.8); ax.set_xlabel("Δ macro-F1 val so với T00")
        extra = "; ".join(f"{r.exp_id}: {r.delta_f1_vs_T00:+.3f}" for _, r in tdf[tdf.exp_id.isin(["T01", "T02"])].iterrows())
        ax.set_title(f"Ablation công thức ({a.backbone}, 1 seed). Trục khởi tạo, ngoài khung: {extra}", fontsize=10)
        ax.grid(alpha=.3, axis="x")
        fig.tight_layout(); fig.savefig(FIGS / "training_ablation.png", dpi=110); plt.close(fig)
    if len(idf) and "lat_p50_ms" in idf:
        d = idf.dropna(subset=["lat_p50_ms"])
        d = d[~d.method.str.contains("fuseBN")]  # gộp BN đo trên ResNet-50/EfficientNet (model khác), xem sheet Latency
        fig, ax = plt.subplots(figsize=(9, 5.5))
        colors = {"I00": "k", "I01": "tab:blue", "I02": "tab:orange", "I04": "tab:green", "I05": "tab:red",
                  "I06": "tab:purple", "I07": "tab:brown", "I08": "tab:gray"}
        for _, r in d.iterrows():
            ax.scatter(r.lat_p50_ms, r.val_macro_f1, color=colors.get(r.exp_id, "c"), s=40, zorder=3)
        for k, (_, r) in enumerate(d.sort_values("lat_p50_ms").iterrows()):
            label = r.method.split(" [")[0].replace("_", " ") + (f" [{r.aggregate}]" if r.aggregate in ("prob", "logit") and r.K_forward > 1 and "crop" not in r.method else "")
            ax.annotate(label, (r.lat_p50_ms, r.val_macro_f1), xytext=(6, -10 + (k % 4) * 7), textcoords="offset points", fontsize=7)
        ax.set_xscale("log"); ax.set_xlabel("độ trễ batch 1 p50 (ms, log), T4 FP32 trừ I08")
        ax.set_ylabel("macro-F1 val"); ax.set_ylim(d.val_macro_f1.min() - 0.002, d.val_macro_f1.max() + 0.002)
        ax.set_title("Suy luận (ConvNeXt-T T00): đánh đổi độ chính xác – độ trễ"); ax.grid(alpha=.3)
        fig.tight_layout(); fig.savefig(FIGS / "inference_tradeoff.png", dpi=110); plt.close(fig)
    for gid, (_, _, test) in per_class.items():
        cm = sum(r["confusion"] for r in test)
        fig, ax = plt.subplots(figsize=(7.5, 6.5))
        ax.imshow(np.log1p(cm), cmap="Blues")
        for i in range(9):
            for j in range(9):
                ax.text(j, i, int(cm[i, j]), ha="center", va="center", fontsize=7,
                        color="white" if cm[i, j] > cm.max() / 3 else "black")
        ax.set_xticks(range(9), D.CLASS_NAMES, rotation=45, ha="right"); ax.set_yticks(range(9), D.CLASS_NAMES)
        ax.set_xlabel("dự đoán"); ax.set_ylabel("nhãn thật")
        ax.set_title(f"Ma trận nhầm lẫn test, {gid}, cộng {len(test)} seed")
        fig.tight_layout(); fig.savefig(FIGS / f"confusion_{gid}.png", dpi=110); plt.close(fig)
    misclassified_grid()
    json.dump({k: {m: list(v) for m, v in st.items()} for k, st in final_stats.items()},
              open(TABLES / "final_stats.json", "w", encoding="utf-8"), indent=2)


def read_table(name: str) -> pd.DataFrame:
    f = TABLES / name
    return pd.read_csv(f) if f.exists() and f.stat().st_size > 2 else pd.DataFrame()


def dataclass_diff(base: dict, cfg: dict) -> dict:
    skip = {"exp_id", "desc", "seed", "save_test_predictions"}
    return {k: v for k, v in cfg.items() if k not in skip and base.get(k) != v}


def combo_desc() -> str:
    f = TABLES / "final_setup.json"
    if f.exists():
        c = json.loads(f.read_text(encoding="utf-8"))["combo"]
        return "công thức {" + ", ".join(f"{k}={v}" for k, v in c.items()) + "}"
    return "công thức kết hợp"


def infer_desc() -> str:
    f = TABLES / "inference_choice.json"
    return json.loads(f.read_text(encoding="utf-8"))["method"] if f.exists() else "suy luận đã chọn"


def misclassified_grid():
    """Ảnh test bị đoán sai giữa Chinee Apple (0) và Snake Weed (7), F01 seed0."""
    import matplotlib.pyplot as plt
    from PIL import Image
    f = SUB / "predictions" / "F01_seed0_test.csv"
    if not f.exists():
        return
    p = read_pred(str(f))
    pairs = [(0, 7), (7, 0)]
    fig, axes = plt.subplots(2, 6, figsize=(14, 5.2))
    for row, (t, pr) in enumerate(pairs):
        idx = np.where((p.y_true == t) & (p.y_pred == pr))[0][:6]
        for j in range(6):
            ax = axes[row, j]; ax.axis("off")
            if j < len(idx):
                i = idx[j]
                ax.imshow(Image.open(DATA / "images" / p.filenames[i]))
                ax.set_title(f"thật {D.CLASS_NAMES[t][:7]} → {D.CLASS_NAMES[pr][:7]}\np={p.probs[i, pr]:.2f}", fontsize=8)
    fig.suptitle("F01 seed0 trên test: nhầm Chinee Apple ↔ Snake Weed")
    fig.tight_layout(); fig.savefig(FIGS / "errors_chinee_snake.png", dpi=90); plt.close(fig)


def style_workbook(path: Path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
            for c in col[1:]:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
        # tô dòng tốt nhất theo macro-F1 val
        hdr = [c.value for c in ws[1]]
        if "val_macro_f1" in hdr and ws.title in ("Backbones", "Training", "Inference"):
            k = hdr.index("val_macro_f1") + 1
            vals = [(ws.cell(r, k).value, r) for r in range(2, ws.max_row + 1)
                    if isinstance(ws.cell(r, k).value, (int, float))]
            if vals:
                best = max(vals)[1]
                for c in ws[best]:
                    c.fill = PatternFill("solid", fgColor="C6EFCE")
        if ws.title == "Summary":
            for c in ws[2]:
                c.fill = PatternFill("solid", fgColor="C6EFCE")
    wb.save(path)


if __name__ == "__main__":
    main()
